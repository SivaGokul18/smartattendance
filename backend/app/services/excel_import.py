import io
import re
import uuid
import random
from typing import Dict, Any, List, Tuple, Optional
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, or_

from app.core.security import get_password_hash
from app.models.user import User, Student, Faculty, RoleEnum
from app.models.academic import Department
from app.models.audit import AuditLog
from app.services.audit_service import log_audit_event

EMAIL_REGEX = re.compile(r"^[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+$")

ROMAN_SEMESTERS = {
    "i": 1, "ii": 2, "iii": 3, "iv": 4, "v": 5,
    "vi": 6, "vii": 7, "viii": 8, "ix": 9, "x": 10
}

CANONICAL_HEADER_KEYWORDS = {
    "email", "mail", "user_email", "email_address",
    "name", "full_name", "student_name", "faculty_name", "staff_name", "fullname",
    "sno", "s_no", "sl_no", "serial_no", "sr_no",
    "dept", "department", "branch",
    "phone", "mobile", "mobile_no", "phone_number", "contact",
    "student_mobile_no", "faculty_mobile_no", "student_mobile", "faculty_mobile",
    "roll_no", "rollno", "roll_number", "roll", "reg_no", "registration_no", "registration_number", "register_no",
    "semester", "sem",
    "section", "sec",
    "mentor", "mentor_email", "mentor_name", "faculty_mentor", "guide",
    "employee_id", "emp_id", "empid", "faculty_id", "staff_id", "staffid",
    "designation", "role_title", "title", "designation_title",
    "assigned_courses", "courses", "subjects",
    "is_mentor",
}

STUDENT_COLUMNS = [
    {"header": "SNO", "key": "sno", "required": False, "example": "1"},
    {"header": "FULL NAME", "key": "full_name", "required": True, "example": "SIVAGOKUL C"},
    {"header": "EMAIL", "key": "email", "required": True, "example": "sivagokulc18@gmail.com"},
    {"header": "ROLL NO", "key": "roll_number", "required": True, "example": "7376242IT303"},
    {"header": "Dept", "key": "department", "required": True, "example": "Information Technology"},
    {"header": "SEMESTER", "key": "semester", "required": True, "example": "V"},
    {"header": "SECTION", "key": "section", "required": True, "example": "SEC A"},
    {"header": "FACULTY NAME", "key": "faculty_name", "required": False, "example": "santhiya"},
    {"header": "STUDENT MOBILE NO", "key": "phone_number", "required": False, "example": "9787380220"},
]

FACULTY_COLUMNS = [
    {"header": "SNO", "key": "sno", "required": False, "example": "1"},
    {"header": "FULL NAME", "key": "full_name", "required": True, "example": "Dr. Santhiya M"},
    {"header": "EMAIL", "key": "email", "required": True, "example": "santhiya.m@campus.edu"},
    {"header": "STAFF ID", "key": "employee_id", "required": True, "example": "STAFF-IT-101"},
    {"header": "Dept", "key": "department", "required": True, "example": "Information Technology"},
    {"header": "DESIGNATION", "key": "designation", "required": False, "example": "Associate Professor"},
    {"header": "FACULTY MOBILE NO", "key": "phone_number", "required": False, "example": "9876501234"},
    {"header": "IS MENTOR", "key": "is_mentor", "required": False, "example": "TRUE"},
]

STUDENT_KEY_ALIASES = {
    "full_name": ["full_name", "name", "student_name", "fullname"],
    "email": ["email", "email_address", "mail", "user_email"],
    "roll_number": ["roll_no", "rollno", "roll_number", "roll", "reg_no", "registration_no", "registration_number", "register_no"],
    "department": ["dept", "department", "branch", "department_name"],
    "semester": ["semester", "sem", "current_semester"],
    "section": ["section", "sec", "current_section", "class_section"],
    "faculty_name": ["faculty_name", "mentor_name", "faculty", "mentor", "mentor_email", "faculty_mentor", "guide"],
    "phone_number": ["student_mobile_no", "student_mobile", "mobile_no", "mobile", "phone", "contact", "phone_number", "student_phone"],
    "emergency_contact": ["emergency_contact", "emergency", "emergency_phone", "guardian_phone"],
    "sno": ["sno", "s_no", "sl_no", "serial_no", "sr_no"],
}

FACULTY_KEY_ALIASES = {
    "full_name": ["full_name", "name", "faculty_name", "staff_name", "fullname"],
    "email": ["email", "email_address", "mail", "user_email"],
    "employee_id": ["staff_id", "employee_id", "emp_id", "empid", "faculty_id", "staffid", "employee_no"],
    "department": ["dept", "department", "branch", "department_name"],
    "designation": ["designation", "title", "role_title", "designation_title"],
    "assigned_courses": ["assigned_courses", "courses", "subjects"],
    "is_mentor": ["is_mentor", "mentor"],
    "phone_number": ["faculty_mobile_no", "faculty_mobile", "mobile_no", "mobile", "phone", "contact", "phone_number", "faculty_phone"],
    "sno": ["sno", "s_no", "sl_no", "serial_no", "sr_no"],
}


def _normalize_header(h: Any) -> str:
    """Cleans up header string by removing requirement marks, trimming, and lowercasing."""
    cleaned = re.sub(r"[\*\(].*?[\)]", "", str(h))
    cleaned = cleaned.replace("*", "").strip().lower()
    cleaned = re.sub(r"\s+", "_", cleaned)
    return cleaned


def parse_semester(val: Any) -> Tuple[int, int]:
    """Parses semester representation (Roman numeral or digit) to (semester_number, year_number)."""
    v_str = str(val).strip().lower()
    clean = re.sub(r"\b(semester|sem)\b", "", v_str).strip()
    if clean in ROMAN_SEMESTERS:
        sem_num = ROMAN_SEMESTERS[clean]
    else:
        digits = re.findall(r"\d+", clean)
        sem_num = int(digits[0]) if digits else 1
    sem_num = max(1, min(10, sem_num))
    year_num = max(1, min(4, (sem_num + 1) // 2))
    return sem_num, year_num


def parse_section(val: Any) -> str:
    """Normalizes section, stripping 'SEC ' or 'SECTION ' prefix (e.g. 'SEC A' -> 'A')."""
    s = str(val).strip()
    m = re.match(r"^(?:sec|section)[\s\-_]*(.+)$", s, re.IGNORECASE)
    if m:
        return m.group(1).strip().upper()
    return s.upper() if s else "A"


def detect_header_and_data(raw_df: pd.DataFrame) -> Tuple[pd.DataFrame, int]:
    """
    Scans the first 15 rows of raw_df to locate the true header row.
    Handles blank top rows, title rows, and institution headers.
    Returns (data_df_with_clean_columns, header_row_index_0_based).
    """
    if raw_df.empty:
        return raw_df, 0

    best_idx = 0
    max_score = -1

    for i in range(min(15, len(raw_df))):
        row_vals = [_normalize_header(c) for c in raw_df.iloc[i].dropna()]
        row_vals = [v for v in row_vals if v]
        score = sum(1 for v in row_vals if v in CANONICAL_HEADER_KEYWORDS)
        if score > max_score:
            max_score = score
            best_idx = i

    if max_score < 2:
        best_idx = 0

    header_values = [_normalize_header(c) for c in raw_df.iloc[best_idx].values]
    data_df = raw_df.iloc[best_idx + 1:].copy().reset_index(drop=True)
    data_df.columns = header_values
    return data_df, best_idx


async def resolve_or_create_department(
    dept_str: str,
    db: AsyncSession,
    all_depts: List[Department]
) -> Department:
    """
    Matches department by code, full name, acronym, or substring.
    Auto-registers the department if not found so imports never fail.
    """
    t_norm = dept_str.strip().lower()
    for d in all_depts:
        if d.code.lower() == t_norm or d.name.lower() == t_norm:
            return d

    words = [w for w in t_norm.split() if w not in {"and", "of", "&", "the"}]
    acronym = "".join(w[0] for w in words).upper()
    if acronym:
        for d in all_depts:
            if d.code.upper() == acronym:
                return d

    for d in all_depts:
        if t_norm in d.name.lower() or d.name.lower() in t_norm:
            return d

    new_code = (acronym[:8] if acronym else t_norm[:5]).upper()
    new_id = f"dept-{new_code.lower()}"
    new_dept = Department(
        id=new_id,
        code=new_code,
        name=dept_str.strip().title(),
        active=True
    )
    db.add(new_dept)
    await db.flush()
    all_depts.append(new_dept)
    return new_dept


async def resolve_or_provision_mentor(
    mentor_input: str,
    default_dept: str,
    db: AsyncSession,
    local_faculty_cache: Dict[str, Faculty],
    all_faculty_users: List[Tuple[User, Faculty]]
) -> Optional[Faculty]:
    """
    Finds existing faculty by email or name.
    If not found, auto-provisions an active faculty member so mentor linkage succeeds.
    """
    mentor_clean = mentor_input.strip()
    if not mentor_clean:
        return None

    m_lower = mentor_clean.lower()
    if m_lower in local_faculty_cache:
        return local_faculty_cache[m_lower]

    for u, f in all_faculty_users:
        if u.email.lower().strip() == m_lower:
            local_faculty_cache[m_lower] = f
            return f
        if u.name.lower().strip() == m_lower or m_lower in u.name.lower().strip():
            local_faculty_cache[m_lower] = f
            return f

    is_email = bool(EMAIL_REGEX.match(mentor_clean))
    slug = re.sub(r"[^a-zA-Z0-9]", "", mentor_clean.lower()) or "mentor"
    email = mentor_clean.lower() if is_email else f"{slug}.faculty@campus.edu"
    name = mentor_clean if not is_email else slug.title()
    employee_id = f"FAC-{slug.upper()[:8]}"

    temp_pw = f"Welcome@{random.randint(1000, 9999)}"
    user_id = f"usr-{uuid.uuid4().hex[:8]}"
    fac_id = f"fac-{uuid.uuid4().hex[:8]}"

    new_user = User(
        id=user_id,
        name=name.title(),
        email=email,
        password_hash=get_password_hash(temp_pw),
        role=RoleEnum.FACULTY,
        department=default_dept,
        must_change_password=True
    )
    new_faculty = Faculty(
        id=fac_id,
        user_id=user_id,
        employee_id=employee_id,
        department=default_dept,
        active=True,
        mentor_group=default_dept
    )
    db.add(new_user)
    db.add(new_faculty)
    await db.flush()

    all_faculty_users.append((new_user, new_faculty))
    local_faculty_cache[m_lower] = new_faculty
    local_faculty_cache[email] = new_faculty
    return new_faculty


def generate_template(import_type: str) -> io.BytesIO:
    """
    Generates a professionally styled .xlsx template for student or faculty bulk import
    matching the institution roster format.
    """
    is_student = import_type.lower() == "student"
    cols = STUDENT_COLUMNS if is_student else FACULTY_COLUMNS

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Import Roster"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    for col_idx, col_info in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=col_idx)
        req_marker = " *" if col_info["required"] else ""
        cell.value = f"{col_info['header']}{req_marker}"
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = align_center
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    sample_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    sample_font = Font(name="Segoe UI", size=10, color="334155")

    # Sample row 1
    for col_idx, col_info in enumerate(cols, start=1):
        cell = ws.cell(row=2, column=col_idx)
        cell.value = col_info["example"]
        cell.font = sample_font
        cell.fill = sample_fill
        cell.alignment = align_left
        cell.border = thin_border
    ws.row_dimensions[2].height = 22

    # Sample row 2
    sample_2_data = []
    if is_student:
        sample_2_data = [
            "2", "SIVANAGU", "sivanagu2006@gmail.com", "7376242IT304",
            "Information Technology", "V", "SEC A", "ramya", "8754124976"
        ]
    else:
        sample_2_data = [
            "2", "Dr. Ramya K", "ramya.k@campus.edu", "STAFF-CS-102",
            "Computer Science and Engineering", "Assistant Professor", "9876505678", "TRUE"
        ]

    for col_idx, val in enumerate(sample_2_data, start=1):
        cell = ws.cell(row=3, column=col_idx)
        cell.value = val
        cell.font = sample_font
        cell.fill = sample_fill
        cell.alignment = align_left
        cell.border = thin_border
    ws.row_dimensions[3].height = 22

    for col_idx, col_info in enumerate(cols, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(col_info["header"]) + 4, len(str(col_info["example"])) + 4, 16)
        ws.column_dimensions[col_letter].width = max_len

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


async def parse_and_validate_excel(
    contents: bytes,
    import_type: str,
    db: AsyncSession
) -> Dict[str, Any]:
    """
    Parses an uploaded Excel or CSV spreadsheet and runs server-side validation against
    database tables. Supports smart header detection (skipping blank/title rows),
    Roman numeral semesters, section normalization, department matching, and mentor resolution.
    """
    is_student = import_type.lower() == "student"
    target_cols = STUDENT_COLUMNS if is_student else FACULTY_COLUMNS
    alias_dict = STUDENT_KEY_ALIASES if is_student else FACULTY_KEY_ALIASES

    # 1. Read raw sheet
    raw_df = None
    try:
        raw_df = pd.read_excel(io.BytesIO(contents), header=None, dtype=str)
    except Exception:
        try:
            raw_df = pd.read_csv(io.BytesIO(contents), header=None, dtype=str)
        except Exception as e:
            raise ValueError(f"Could not read spreadsheet file: {str(e)}")

    if raw_df is None or raw_df.empty:
        return {
            "import_type": import_type,
            "total_rows": 0,
            "valid_rows_count": 0,
            "warning_rows_count": 0,
            "error_rows_count": 0,
            "preview_rows": [],
            "headers": [c["header"] for c in target_cols],
        }

    # 2. Smart header detection
    df, header_row_idx = detect_header_and_data(raw_df)
    df_cols_normalized = list(df.columns)

    # 3. Rename columns using aliases
    for canonical_key, aliases in alias_dict.items():
        if canonical_key not in df.columns:
            matched_alias = next((a for a in aliases if a in df_cols_normalized), None)
            if matched_alias:
                df.rename(columns={matched_alias: canonical_key}, inplace=True)
                df_cols_normalized = list(df.columns)

    # Verify required headers
    missing_required_headers = []
    for c in target_cols:
        key = c["key"]
        if c["required"] and key not in df.columns:
            missing_required_headers.append(c["header"])

    if missing_required_headers:
        raise ValueError(
            f"Spreadsheet is missing required column headers: {', '.join(missing_required_headers)}. "
            f"Please verify headers match the template."
        )

    # 4. Reference data lookup
    all_users = (await db.execute(select(User))).scalars().all()
    user_by_email = {u.email.lower().strip(): u for u in all_users if u.email}
    user_by_id = {u.id: u for u in all_users}

    student_by_roll: Dict[str, Student] = {}
    faculty_by_emp: Dict[str, Faculty] = {}
    if is_student:
        all_students = (await db.execute(select(Student))).scalars().all()
        student_by_roll = {s.roll_number.lower().strip(): s for s in all_students if s.roll_number}
    else:
        all_faculty = (await db.execute(select(Faculty))).scalars().all()
        faculty_by_emp = {f.employee_id.lower().strip(): f for f in all_faculty if f.employee_id}

    all_depts = (await db.execute(select(Department))).scalars().all()
    all_depts_list = list(all_depts)

    faculty_users_raw = (
        await db.execute(
            select(User, Faculty).join(Faculty, Faculty.user_id == User.id)
        )
    ).all()
    all_faculty_users = list(faculty_users_raw)
    faculty_email_or_name_set = {
        u.email.lower().strip() for u, f in all_faculty_users if u.email
    } | {
        u.name.lower().strip() for u, f in all_faculty_users if u.name
    }

    seen_file_emails: Dict[str, int] = {}
    seen_file_ids: Dict[str, int] = {}

    preview_rows = []
    valid_count = 0
    warning_count = 0
    error_count = 0

    # 5. Iterate each row
    for row_idx, row in df.iterrows():
        excel_row_num = header_row_idx + 2 + int(row_idx)

        row_values = [str(v).strip() for v in row.values if pd.notna(v) and str(v).strip() != ""]
        if not row_values:
            continue

        row_data = {}
        row_errors = []
        row_warnings = []

        for c in target_cols:
            k = c["key"]
            val = str(row.get(k, "")).strip() if pd.notna(row.get(k)) else ""
            if val.lower() == "nan":
                val = ""
            row_data[k] = val

        # Validate Full Name
        name = row_data.get("full_name", "").strip()
        if not name:
            row_errors.append("Full Name is required.")

        # Validate Email
        email = row_data.get("email", "").strip().lower()
        if not email:
            row_errors.append("Email is required.")
        elif not EMAIL_REGEX.match(email):
            row_errors.append(f"Invalid email address format: '{email}'.")
        else:
            if email in seen_file_emails:
                row_errors.append(f"Duplicate email '{email}' in row {seen_file_emails[email]} and row {excel_row_num}.")
            else:
                seen_file_emails[email] = excel_row_num

        # Match existing user (by email or by identifier)
        matched_user: Optional[User] = None
        action = "create"

        if email and email in user_by_email:
            matched_user = user_by_email[email]

        # Validate Roll Number / Employee ID
        if is_student:
            roll = row_data.get("roll_number", "").strip().upper()
            if not roll:
                row_errors.append("Roll Number is required.")
            else:
                r_lower = roll.lower()
                if r_lower in seen_file_ids:
                    row_errors.append(f"Duplicate Roll Number '{roll}' in row {seen_file_ids[r_lower]} and row {excel_row_num}.")
                else:
                    seen_file_ids[r_lower] = excel_row_num

                if not matched_user and r_lower in student_by_roll:
                    s_owner = student_by_roll[r_lower]
                    matched_user = user_by_id.get(s_owner.user_id)

                if r_lower in student_by_roll:
                    s_owner = student_by_roll[r_lower]
                    if matched_user and matched_user.id != s_owner.user_id:
                        row_errors.append(f"Roll Number '{roll}' is already assigned to another student.")
        else:
            emp = row_data.get("employee_id", "").strip().upper()
            if not emp:
                row_errors.append("Employee ID is required.")
            else:
                e_lower = emp.lower()
                if e_lower in seen_file_ids:
                    row_errors.append(f"Duplicate Employee ID '{emp}' in row {seen_file_ids[e_lower]} and row {excel_row_num}.")
                else:
                    seen_file_ids[e_lower] = excel_row_num

                if not matched_user and e_lower in faculty_by_emp:
                    f_owner = faculty_by_emp[e_lower]
                    matched_user = user_by_id.get(f_owner.user_id)

                if e_lower in faculty_by_emp:
                    f_owner = faculty_by_emp[e_lower]
                    if matched_user and matched_user.id != f_owner.user_id:
                        row_errors.append(f"Employee ID '{emp}' is already assigned to another faculty member.")

        if matched_user:
            expected_role = RoleEnum.STUDENT if is_student else RoleEnum.FACULTY
            if matched_user.role != expected_role:
                row_errors.append(f"Email/Identifier matches existing user with role '{matched_user.role}', not '{expected_role}'.")
            else:
                action = "update"
                row_warnings.append(f"Existing {import_type} record will be updated.")

        # Validate & Match Department
        dept_str = row_data.get("department", "").strip()
        if not dept_str:
            row_errors.append("Department is required.")
        else:
            t_norm = dept_str.lower()
            matched_dept = next(
                (d for d in all_depts_list if d.code.lower() == t_norm or d.name.lower() == t_norm),
                None
            )
            if not matched_dept:
                words = [w for w in t_norm.split() if w not in {"and", "of", "&", "the"}]
                acronym = "".join(w[0] for w in words).upper()
                if acronym:
                    matched_dept = next((d for d in all_depts_list if d.code.upper() == acronym), None)
            if not matched_dept:
                matched_dept = next((d for d in all_depts_list if t_norm in d.name.lower() or d.name.lower() in t_norm), None)

            if matched_dept:
                row_data["department"] = matched_dept.name
            else:
                row_warnings.append(f"New department '{dept_str}' will be automatically created.")

        # Student specific formatting
        if is_student:
            sem_raw = row_data.get("semester", "").strip()
            if not sem_raw:
                row_errors.append("Semester is required (e.g. V or 5).")
            else:
                sem_num, year_num = parse_semester(sem_raw)
                row_data["semester"] = str(sem_num)
                row_data["year"] = year_num

            sec_raw = row_data.get("section", "").strip()
            if not sec_raw:
                row_errors.append("Section is required (e.g. SEC A or A).")
            else:
                row_data["section"] = parse_section(sec_raw)

            faculty_mentor_str = row_data.get("faculty_name", "").strip()
            if faculty_mentor_str:
                fm_lower = faculty_mentor_str.lower()
                is_known = any(
                    fm_lower in name_or_email for name_or_email in faculty_email_or_name_set
                )
                if not is_known:
                    row_warnings.append(
                        f"Mentor account for '{faculty_mentor_str}' will be auto-provisioned upon import."
                    )

        if row_errors:
            status = "error"
            error_count += 1
            all_messages = row_errors + row_warnings
        elif row_warnings:
            status = "warning"
            warning_count += 1
            all_messages = row_warnings
        else:
            status = "valid"
            valid_count += 1
            all_messages = []

        preview_rows.append({
            "row_number": excel_row_num,
            "data": row_data,
            "action": action,
            "status": status,
            "errors": all_messages,
        })

    return {
        "import_type": import_type,
        "total_rows": len(preview_rows),
        "valid_rows_count": valid_count,
        "warning_rows_count": warning_count,
        "error_rows_count": error_count,
        "preview_rows": preview_rows,
        "headers": [c["header"] for c in target_cols],
    }


async def process_bulk_import(
    import_type: str,
    rows: List[Dict[str, Any]],
    admin_user: User,
    db: AsyncSession
) -> Dict[str, Any]:
    """
    Executes bulk import:
    Creates new User accounts or updates existing records, resolves/provisions departments & mentors,
    updates/creates Student/Faculty profiles, sets must_change_password=True for new accounts,
    logs audit events, and returns generated credentials and update summary.
    """
    is_student = import_type.lower() == "student"

    all_users = (await db.execute(select(User))).scalars().all()
    user_by_email = {u.email.lower().strip(): u for u in all_users if u.email}
    user_by_id = {u.id: u for u in all_users}

    student_by_roll: Dict[str, Student] = {}
    faculty_by_emp: Dict[str, Faculty] = {}
    student_by_user_id: Dict[str, Student] = {}
    faculty_by_user_id: Dict[str, Faculty] = {}

    if is_student:
        all_students = (await db.execute(select(Student))).scalars().all()
        for s in all_students:
            if s.roll_number:
                student_by_roll[s.roll_number.lower().strip()] = s
            student_by_user_id[s.user_id] = s
    else:
        all_faculty = (await db.execute(select(Faculty))).scalars().all()
        for f in all_faculty:
            if f.employee_id:
                faculty_by_emp[f.employee_id.lower().strip()] = f
            faculty_by_user_id[f.user_id] = f

    all_depts_list = list((await db.execute(select(Department))).scalars().all())

    faculty_users_raw = (
        await db.execute(
            select(User, Faculty).join(Faculty, Faculty.user_id == User.id)
        )
    ).all()
    all_faculty_users = list(faculty_users_raw)
    local_faculty_cache: Dict[str, Faculty] = {}

    imported_count = 0
    updated_count = 0
    skipped_count = 0
    skipped_errors = []
    generated_credentials = []

    for row_item in rows:
        row_num = row_item.get("row_number", "?")
        data = row_item.get("data", {})
        status = row_item.get("status", "valid")

        if status == "error":
            skipped_count += 1
            skipped_errors.append({
                "row_number": row_num,
                "data": data,
                "reason": "; ".join(row_item.get("errors", ["Validation error"]))
            })
            continue

        email = data.get("email", "").strip().lower()
        name = data.get("full_name", "").strip()
        raw_dept = data.get("department", "Information Technology").strip()
        phone = data.get("phone_number", "").strip() or None

        # Resolve or auto-create department
        dept_obj = await resolve_or_create_department(raw_dept, db, all_depts_list)
        dept_name = dept_obj.name

        # Check existing user
        matched_user: Optional[User] = None
        if email and email in user_by_email:
            matched_user = user_by_email[email]
        elif is_student:
            r_str = data.get("roll_number", "").strip().lower()
            if r_str in student_by_roll:
                matched_user = user_by_id.get(student_by_roll[r_str].user_id)
        else:
            e_str = data.get("employee_id", "").strip().lower()
            if e_str in faculty_by_emp:
                matched_user = user_by_id.get(faculty_by_emp[e_str].user_id)

        if matched_user:
            # Update existing user & profile
            if name:
                matched_user.name = name
            matched_user.department = dept_name
            if phone is not None:
                matched_user.phone = phone
            if email and matched_user.email != email:
                matched_user.email = email
                user_by_email[email] = matched_user

            if is_student:
                stu = student_by_user_id.get(matched_user.id)
                if stu:
                    roll = data.get("roll_number", "").strip().upper()
                    if roll:
                        stu.roll_number = roll
                    stu.department = dept_name

                    sec_raw = data.get("section", "").strip()
                    if sec_raw:
                        stu.section = parse_section(sec_raw)

                    sem_raw = data.get("semester", "")
                    if sem_raw:
                        sem_num, year_num = parse_semester(sem_raw)
                        stu.year = year_num

                    mentor_input = data.get("faculty_name", "").strip()
                    if mentor_input:
                        mentor_fac = await resolve_or_provision_mentor(
                            mentor_input=mentor_input,
                            default_dept=dept_name,
                            db=db,
                            local_faculty_cache=local_faculty_cache,
                            all_faculty_users=all_faculty_users
                        )
                        if mentor_fac:
                            stu.mentor_id = mentor_fac.id
            else:
                fac = faculty_by_user_id.get(matched_user.id)
                if fac:
                    emp = data.get("employee_id", "").strip().upper()
                    if emp:
                        fac.employee_id = emp
                    fac.department = dept_name
                    desig = data.get("designation", "").strip()
                    if desig:
                        fac.designation = desig
                    is_mentor_str = str(data.get("is_mentor", "")).strip().lower()
                    is_mentor = is_mentor_str in ["true", "1", "yes", "y"]
                    fac.is_mentor = is_mentor
                    fac.mentor_group = dept_name if is_mentor else None

            updated_count += 1
        else:
            # Create new user
            temp_password = f"Welcome@{random.randint(1000, 9999)}"
            hashed_password = get_password_hash(temp_password)

            user_role = RoleEnum.STUDENT if is_student else RoleEnum.FACULTY
            user_id = f"usr-{uuid.uuid4().hex[:8]}"

            new_user = User(
                id=user_id,
                name=name,
                email=email,
                password_hash=hashed_password,
                role=user_role,
                department=dept_name,
                phone=phone,
                must_change_password=True
            )
            db.add(new_user)
            user_by_email[email] = new_user
            user_by_id[user_id] = new_user

            identifier = ""
            if is_student:
                roll = data.get("roll_number", "").strip().upper()
                identifier = roll

                sem_num, year_num = parse_semester(data.get("semester", "5"))
                sec_str = parse_section(data.get("section", "A"))

                mentor_input = data.get("faculty_name", "").strip()
                mentor_faculty = None
                if mentor_input:
                    mentor_faculty = await resolve_or_provision_mentor(
                        mentor_input=mentor_input,
                        default_dept=dept_name,
                        db=db,
                        local_faculty_cache=local_faculty_cache,
                        all_faculty_users=all_faculty_users
                    )

                new_student = Student(
                    id=f"stu-{uuid.uuid4().hex[:8]}",
                    user_id=new_user.id,
                    roll_number=roll,
                    department=dept_name,
                    year=year_num,
                    section=sec_str,
                    attendance_rate=90.0,
                    face_id_status="pending",
                    mentor_id=mentor_faculty.id if mentor_faculty else None
                )
                db.add(new_student)
                student_by_user_id[new_user.id] = new_student
                if roll:
                    student_by_roll[roll.lower()] = new_student
            else:
                emp = data.get("employee_id", "").strip().upper()
                identifier = emp

                desig = data.get("designation", "").strip() or "Faculty"
                is_mentor_str = str(data.get("is_mentor", "")).strip().lower()
                is_mentor = is_mentor_str in ["true", "1", "yes", "y"]

                new_faculty = Faculty(
                    id=f"fac-{uuid.uuid4().hex[:8]}",
                    user_id=new_user.id,
                    employee_id=emp,
                    department=dept_name,
                    designation=desig,
                    active=True,
                    is_mentor=is_mentor,
                    mentor_group=dept_name if is_mentor else None
                )
                db.add(new_faculty)
                faculty_by_user_id[new_user.id] = new_faculty
                if emp:
                    faculty_by_emp[emp.lower()] = new_faculty

            imported_count += 1
            generated_credentials.append({
                "name": name,
                "email": email,
                "identifier": identifier,
                "role": user_role,
                "temporaryPassword": temp_password,
                "department": dept_name
            })

    await log_audit_event(
        db=db,
        actor_id=admin_user.id,
        actor_role=admin_user.role,
        action_type="bulk_import",
        target_entity=f"bulk_{import_type}",
        metadata={
            "import_type": import_type,
            "imported_count": imported_count,
            "updated_count": updated_count,
            "skipped_count": skipped_count,
            "total_processed": len(rows),
        }
    )

    await db.commit()

    message = (
        f"Successfully processed {import_type} import: {imported_count} created, "
        f"{updated_count} updated, {skipped_count} skipped."
    )

    return {
        "imported": imported_count,
        "updated": updated_count,
        "skipped": skipped_count,
        "errors": skipped_errors,
        "credentials": generated_credentials,
        "message": message,
    }
