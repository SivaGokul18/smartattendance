import asyncio
import os
import re
import uuid
import openpyxl
from sqlalchemy import select
from app.core.config import settings
from app.core.database import AsyncSessionLocal, engine, Base
from app.core.security import get_password_hash
from app.models.user import User, Student, Faculty, RoleEnum
from app.models.academic import Department, ClassSection
from app.models.settings import AdminSettings


ROMAN_SEMESTERS = {
    "i": 1, "ii": 2, "iii": 3, "iv": 4, "v": 5,
    "vi": 6, "vii": 7, "viii": 8, "ix": 9, "x": 10
}


def parse_semester_to_year(sem_val) -> int:
    if not sem_val:
        return 3
    s = str(sem_val).strip().lower()
    sem_num = ROMAN_SEMESTERS.get(s)
    if sem_num is None:
        try:
            sem_num = int(s)
        except ValueError:
            sem_num = 5
    # Year = ceil(sem / 2)
    return max(1, min(4, (sem_num + 1) // 2))


def clean_section(sec_val) -> str:
    if not sec_val:
        return "A"
    s = str(sec_val).strip().upper()
    s = re.sub(r"^SEC(TION)?\s*", "", s).strip()
    return s if s else "A"


async def load_roster_from_excel(file_path: str = "institutional_roster.xlsx"):
    full_path = os.path.abspath(file_path)
    if not os.path.exists(full_path):
        # Check parent dir
        parent_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", file_path))
        if os.path.exists(parent_path):
            full_path = parent_path
        else:
            raise FileNotFoundError(f"Excel file not found at: {full_path}")

    print(f"Loading roster from: {full_path}")
    wb = openpyxl.load_workbook(full_path, data_only=True)

    async with AsyncSessionLocal() as db:
        # 1. Ensure Super Admin exists
        existing_admin = (await db.execute(select(User).where(User.email == "admin@campus.edu"))).scalar_one_or_none()
        if not existing_admin:
            admin_user = User(
                id="usr-admin-1",
                name="Super Admin Administrator",
                email="admin@campus.edu",
                password_hash=get_password_hash("admin123"),
                role=RoleEnum.ADMIN,
                department="Institutional Administration",
                phone="+91 98765 00001"
            )
            db.add(admin_user)
            print("  + [ADMIN] Initialized admin@campus.edu")

        # 2. Ensure default institutional settings
        existing_settings = (await db.execute(select(AdminSettings))).scalar_one_or_none()
        if not existing_settings:
            db.add(AdminSettings(
                id="institution-settings-default",
                ble_signal_range=15.0,
                ble_rssi_threshold=-75.0,
                face_confidence_threshold=85.0,
                liveness_check_enabled=True,
                auto_sync_offline=True,
                faculty_manual_override_allowed=True
            ))
            print("  + [SETTINGS] Initialized default settings")

        # 3. Load Departments Cache
        all_depts = list((await db.execute(select(Department))).scalars().all())
        dept_by_name = {d.name.lower().strip(): d for d in all_depts}

        async def get_or_create_department(dept_name: str) -> str:
            name_clean = dept_name.strip()
            key = name_clean.lower()
            if key in dept_by_name:
                return dept_by_name[key].name
            
            # Generate code
            words = [w for w in re.split(r"[\s&,-]+", name_clean) if w]
            code = "".join(w[0].upper() for w in words if w.lower() not in ["and", "of", "in", "the"])
            if not code or len(code) < 2:
                code = name_clean[:4].upper()
            
            new_dept = Department(
                id=f"dept-{uuid.uuid4().hex[:6]}",
                code=code[:10],
                name=name_clean,
                active=True
            )
            db.add(new_dept)
            dept_by_name[key] = new_dept
            return name_clean

        # 4. Load Faculty Sheet
        faculty_cache = {}
        if "Faculty" in wb.sheetnames:
            ws_fac = wb["Faculty"]
            rows = list(ws_fac.iter_rows(values_only=True))
            if len(rows) > 1:
                headers = [str(h).strip().upper() if h else "" for h in rows[0]]
                name_idx = headers.index("FULL NAME") if "FULL NAME" in headers else 1
                email_idx = headers.index("EMAIL") if "EMAIL" in headers else 2
                emp_idx = headers.index("STAFF ID") if "STAFF ID" in headers else 3
                dept_idx = headers.index("DEPT") if "DEPT" in headers else 4
                desig_idx = headers.index("DESIGNATION") if "DESIGNATION" in headers else 5
                phone_idx = headers.index("FACULTY MOBILE NO") if "FACULTY MOBILE NO" in headers else 6
                is_mentor_idx = headers.index("IS MENTOR") if "IS MENTOR" in headers else 7
                pwd_idx = headers.index("PASSWORD") if "PASSWORD" in headers else -1

                for row in rows[1:]:
                    if not row or not any(row):
                        continue
                    name = str(row[name_idx] or "").strip()
                    email = str(row[email_idx] or "").strip().lower()
                    emp_id = str(row[emp_idx] or "").strip()
                    dept_raw = str(row[dept_idx] or "Information Technology").strip()
                    desig = str(row[desig_idx] or "Faculty").strip() if desig_idx < len(row) and row[desig_idx] else "Faculty"
                    phone = str(row[phone_idx] or "").strip() if phone_idx < len(row) and row[phone_idx] else None
                    is_mentor_val = str(row[is_mentor_idx] or "FALSE").strip().upper() if is_mentor_idx < len(row) and row[is_mentor_idx] else "FALSE"
                    is_mentor_bool = is_mentor_val in ["TRUE", "YES", "1"]
                    pwd = str(row[pwd_idx] or "faculty123").strip() if pwd_idx != -1 and pwd_idx < len(row) and row[pwd_idx] else "faculty123"

                    if not email or not name:
                        continue

                    dept_name = await get_or_create_department(dept_raw)

                    # Check existing user
                    user = (await db.execute(select(User).where(User.email == email))).scalar_one_or_none()
                    if not user and emp_id:
                        fac_exist = (await db.execute(select(Faculty).where(Faculty.employee_id == emp_id))).scalar_one_or_none()
                        if fac_exist:
                            user = (await db.execute(select(User).where(User.id == fac_exist.user_id))).scalar_one_or_none()
                            if user and email:
                                user.email = email

                    if not user:
                        user = User(
                            id=f"usr-fac-{uuid.uuid4().hex[:8]}",
                            name=name,
                            email=email,
                            password_hash=get_password_hash(pwd),
                            role=RoleEnum.FACULTY,
                            department=dept_name,
                            phone=phone,
                            photo_url="https://images.unsplash.com/photo-1534528741775-53994a69daeb?w=150"
                        )
                        db.add(user)
                        await db.flush()
                    else:
                        user.name = name
                        user.department = dept_name
                        user.phone = phone
                        user.password_hash = get_password_hash(pwd)

                    fac = (await db.execute(select(Faculty).where(Faculty.user_id == user.id))).scalar_one_or_none()
                    if not fac:
                        fac = Faculty(
                            id=f"fac-{uuid.uuid4().hex[:8]}",
                            user_id=user.id,
                            employee_id=emp_id or f"STAFF-{uuid.uuid4().hex[:4].upper()}",
                            department=dept_name,
                            designation=desig,
                            active=True,
                            is_mentor=is_mentor_bool,
                            mentor_group="Section A" if is_mentor_bool else None
                        )
                        db.add(fac)
                        await db.flush()
                    else:
                        fac.employee_id = emp_id or fac.employee_id
                        fac.department = dept_name
                        fac.designation = desig
                        fac.is_mentor = is_mentor_bool
                        fac.mentor_group = "Section A" if is_mentor_bool else None

                    faculty_cache[name.lower()] = fac
                    faculty_cache[email] = fac
                    print(f"  + [FACULTY] {name} ({email}) - {desig} - Employee ID: {fac.employee_id} - Mentor: {is_mentor_bool}")

        # 5. Load Students Sheet
        if "Students" in wb.sheetnames:
            ws_stu = wb["Students"]
            rows = list(ws_stu.iter_rows(values_only=True))
            if len(rows) > 1:
                headers = [str(h).strip().upper() if h else "" for h in rows[0]]
                name_idx = headers.index("FULL NAME") if "FULL NAME" in headers else 1
                email_idx = headers.index("EMAIL") if "EMAIL" in headers else 2
                roll_idx = headers.index("ROLL NO") if "ROLL NO" in headers else 3
                dept_idx = headers.index("DEPT") if "DEPT" in headers else 4
                sem_idx = headers.index("SEMESTER") if "SEMESTER" in headers else 5
                sec_idx = headers.index("SECTION") if "SECTION" in headers else 6
                mentor_idx = headers.index("FACULTY NAME") if "FACULTY NAME" in headers else 7
                phone_idx = headers.index("STUDENT MOBILE NO") if "STUDENT MOBILE NO" in headers else 8
                pwd_idx = headers.index("PASSWORD") if "PASSWORD" in headers else -1

                for row in rows[1:]:
                    if not row or not any(row):
                        continue
                    name = str(row[name_idx] or "").strip()
                    email = str(row[email_idx] or "").strip().lower()
                    roll = str(row[roll_idx] or "").strip()
                    dept_raw = str(row[dept_idx] or "Information Technology").strip()
                    sem_val = row[sem_idx] if sem_idx < len(row) else "V"
                    sec_val = row[sec_idx] if sec_idx < len(row) else "SEC A"
                    mentor_name = str(row[mentor_idx] or "").strip() if mentor_idx < len(row) and row[mentor_idx] else ""
                    phone = str(row[phone_idx] or "").strip() if phone_idx < len(row) and row[phone_idx] else None
                    pwd = str(row[pwd_idx] or "student123").strip() if pwd_idx != -1 and pwd_idx < len(row) and row[pwd_idx] else "student123"

                    if not email or not name:
                        continue

                    dept_name = await get_or_create_department(dept_raw)
                    year = parse_semester_to_year(sem_val)
                    section = clean_section(sec_val)

                    # Match Mentor
                    mentor_id = None
                    if mentor_name:
                        matched_fac = faculty_cache.get(mentor_name.lower())
                        if not matched_fac:
                            # Fuzzy match
                            for k, f in faculty_cache.items():
                                if mentor_name.lower() in k:
                                    matched_fac = f
                                    break
                        if matched_fac:
                            mentor_id = matched_fac.id

                    # Check or create ClassSection
                    sec_title = f"{dept_name} - Year {year} - Section {section}"
                    cls_sec = (await db.execute(select(ClassSection).where(ClassSection.name == sec_title))).scalar_one_or_none()
                    if not cls_sec:
                        cls_sec = ClassSection(
                            id=f"sec-{uuid.uuid4().hex[:8]}",
                            name=sec_title,
                            department=dept_name,
                            year=year,
                            section=section,
                            student_count=60
                        )
                        db.add(cls_sec)

                    user = (await db.execute(select(User).where(User.email == email))).scalar_one_or_none()
                    if not user and roll:
                        stu_exist = (await db.execute(select(Student).where(Student.roll_number == roll))).scalar_one_or_none()
                        if stu_exist:
                            user = (await db.execute(select(User).where(User.id == stu_exist.user_id))).scalar_one_or_none()
                            if user and email:
                                user.email = email

                    if not user:
                        user = User(
                            id=f"usr-stu-{uuid.uuid4().hex[:8]}",
                            name=name,
                            email=email,
                            password_hash=get_password_hash(pwd),
                            role=RoleEnum.STUDENT,
                            department=dept_name,
                            phone=phone,
                            photo_url="https://images.unsplash.com/photo-1539571696357-5a69c17a67c6?w=150"
                        )
                        db.add(user)
                        await db.flush()
                    else:
                        user.name = name
                        user.department = dept_name
                        user.phone = phone
                        user.password_hash = get_password_hash(pwd)

                    stu = (await db.execute(select(Student).where(Student.user_id == user.id))).scalar_one_or_none()
                    if not stu:
                        stu = Student(
                            id=f"stu-{uuid.uuid4().hex[:8]}",
                            user_id=user.id,
                            roll_number=roll or f"2026{uuid.uuid4().hex[:4].upper()}",
                            department=dept_name,
                            year=year,
                            section=section,
                            attendance_rate=94.5,
                            face_id_status="enrolled",
                            mentor_id=mentor_id
                        )
                        db.add(stu)
                    else:
                        stu.roll_number = roll or stu.roll_number
                        stu.department = dept_name
                        stu.year = year
                        stu.section = section
                        if mentor_id:
                            stu.mentor_id = mentor_id

                    print(f"  + [STUDENT] {name} ({email}) - Roll: {stu.roll_number} - Section: {section}")

        await db.commit()
        print("\n[SUCCESS] Excel Roster loaded into MySQL database successfully!")


if __name__ == "__main__":
    asyncio.run(load_roster_from_excel())
